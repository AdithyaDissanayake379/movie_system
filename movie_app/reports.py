"""Excel reports and rental frequency chart."""
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from movie_app.db import fetch_all

REPORT_QUERIES = {
 "Currently Rented": """SELECT r.rental_id,m.title,c.name,r.rental_date,r.due_date
 FROM rentals r JOIN movies m ON m.movie_id=r.movie_id
 JOIN customers c ON c.customer_id=r.customer_id WHERE r.return_date IS NULL""",
 "Overdue Rentals": """SELECT r.rental_id,m.title,c.name,r.rental_date,r.due_date
 FROM rentals r JOIN movies m ON m.movie_id=r.movie_id
 JOIN customers c ON c.customer_id=r.customer_id
 WHERE r.return_date IS NULL AND r.due_date < CURDATE()""",
 "By Genre": """SELECT m.genre,COUNT(*) AS rental_count
 FROM rentals r JOIN movies m ON m.movie_id=r.movie_id GROUP BY m.genre ORDER BY rental_count DESC""",
 "By Producer": """SELECT m.producer,COUNT(*) AS rental_count
 FROM rentals r JOIN movies m ON m.movie_id=r.movie_id GROUP BY m.producer ORDER BY rental_count DESC"""
}

def export_excel(destination=None):
    destination = Path(destination or f"reports/movie_report_{date.today().isoformat()}.xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, query in REPORT_QUERIES.items():
        ws = wb.create_sheet(name)
        rows = fetch_all(query)
        if not rows:
            ws.append(["No records"])
            continue
        ws.append(list(rows[0].keys()))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24476A")
        for row in rows:
            ws.append(list(row.values()))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = min(40, max(len(str(c.value or "")) for c in column) + 3)
            ws.column_dimensions[column[0].column_letter].width = width
    wb.save(destination)
    return destination

def show_chart():
    import matplotlib.pyplot as plt
    rows = fetch_all(REPORT_QUERIES["By Genre"])
    if not rows:
        raise ValueError("There are no rentals to visualize yet")
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar([r["genre"] for r in rows], [r["rental_count"] for r in rows])
    ax.set_title("Rentals by Genre")
    ax.set_xlabel("Genre")
    ax.set_ylabel("Number of rentals")
    fig.tight_layout()
    plt.show()